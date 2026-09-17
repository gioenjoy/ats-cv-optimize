#!/usr/bin/env python3
"""Renderizza un CV italiano ATS-safe in formato DOCX a partire da un JSON.

Il documento prodotto è deliberatamente sobrio: una sola colonna, nessuna tabella,
nessuna casella di testo, nessuna immagine, contatti nel corpo del documento.
Ogni elemento grafico in più è un punto di rottura per i parser ATS.

Uso:
    python3 render_docx.py --input cv.json --output CV_Nome_Cognome.docx
    python3 render_docx.py --input cv.json --check      # solo validazione
"""

import argparse
import json
import re
import sys

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    sys.exit(
        "Manca la libreria python-docx.\n"
        "Installala con:  pip3 install python-docx"
    )

# --- Impostazioni tipografiche -------------------------------------------------

FONT = "Calibri"
CORPO_PT = 10.5
NOME_PT = 20
TITOLO_SEZIONE_PT = 11
CONTATTI_PT = 9.5
PRIVACY_PT = 8
MARGINE_CM = 1.6
GRIGIO = RGBColor(0x40, 0x40, 0x40)
NERO = RGBColor(0x00, 0x00, 0x00)

ORDINE_PREDEFINITO = {
    "neolaureato": ["profilo", "formazione", "esperienze", "progetti",
                    "competenze", "certificazioni", "lingue", "altro"],
    "junior": ["profilo", "esperienze", "competenze", "formazione",
               "certificazioni", "lingue", "progetti", "altro"],
    "mid": ["profilo", "competenze", "esperienze", "formazione",
            "certificazioni", "lingue", "progetti", "altro"],
    "senior": ["profilo", "competenze", "esperienze", "formazione",
               "certificazioni", "lingue", "progetti", "altro"],
    "manageriale": ["profilo", "esperienze", "competenze", "formazione",
                    "certificazioni", "lingue", "altro"],
    "cambio_settore": ["profilo", "competenze", "esperienze", "formazione",
                       "certificazioni", "lingue", "progetti", "altro"],
}

TITOLI_SEZIONE = {
    "profilo": "Profilo professionale",
    "competenze": "Competenze",
    "esperienze": "Esperienza professionale",
    "formazione": "Formazione",
    "certificazioni": "Certificazioni",
    "lingue": "Lingue",
    "progetti": "Progetti",
    "altro": "Altre attività",
}

PRIVACY_PREDEFINITA = (
    "Autorizzo il trattamento dei miei dati personali contenuti nel presente "
    "curriculum vitae ai sensi del Regolamento (UE) 2016/679 (GDPR) e del "
    "D.lgs. 196/2003 come modificato dal D.lgs. 101/2018."
)

MESI_IT = {"gen", "feb", "mar", "apr", "mag", "giu",
           "lug", "ago", "set", "ott", "nov", "dic"}


# --- Validazione ---------------------------------------------------------------

APERTURE_VIETATE = re.compile(
    r"^\s*(ho\b|mi sono\b|sono stat|abbiamo\b|avevo\b|ero\b|"
    r"responsabile d[ei]\b|mi occupav)", re.IGNORECASE)

# Participi accordati più frequenti in apertura di bullet.
PARTICIPI_ACCORDATI = re.compile(
    r"^\s*(progettat|sviluppat|realizzat|gestit|coordinat|ridott|aumentat|"
    r"analizzat|implementat|automatizzat|creat|curat|seguit|organizzat|"
    r"ottimizzat|monitorat|elaborat|supervisionat|negoziat|acquisit)[aie]\b",
    re.IGNORECASE)

DATA_VALIDA = re.compile(
    r"^(gen|feb|mar|apr|mag|giu|lug|ago|set|ott|nov|dic)\s+\d{4}$")

EMOJI = re.compile(
    "[" "\U0001F300-\U0001FAFF" "\U00002600-\U000027BF"
    "\U0001F1E6-\U0001F1FF" "\U00002B00-\U00002BFF" "]")

DECORATIVI = re.compile(r"[█░▒▓■●★☆➤✦]")

CF = re.compile(r"\b[A-Z]{6}\d{2}[A-EHLMPR-T]\d{2}[A-Z]\d{3}[A-Z]\b")


def valida(cv):
    """Restituisce (errori, avvisi)."""
    errori, avvisi = [], []

    intest = cv.get("intestazione") or {}
    if not intest.get("nome") or not intest.get("cognome"):
        errori.append("intestazione.nome e intestazione.cognome sono obbligatori")

    testo_completo = json.dumps(cv, ensure_ascii=False)
    if EMOJI.search(testo_completo):
        errori.append("il CV contiene emoji: rimuovile")
    if DECORATIVI.search(testo_completo):
        errori.append("il CV contiene caratteri decorativi o barre di valutazione")
    if CF.search(testo_completo):
        errori.append("sembra presente un codice fiscale: non va mai in un CV")

    citta = intest.get("citta") or ""
    if re.search(r"\b(via|viale|corso|piazza|piazzale|largo|vicolo)\b",
                 citta, re.IGNORECASE):
        avvisi.append(
            "intestazione.citta contiene un indirizzo: bastano città e provincia")

    for esp in cv.get("esperienze") or []:
        eti = f"{esp.get('ruolo', '?')} @ {esp.get('azienda', '?')}"
        for campo in ("dal", "al"):
            v = (esp.get(campo) or "").strip()
            if not v:
                continue
            if campo == "al" and v.lower() in ("oggi", "in corso", "presente"):
                continue
            if not DATA_VALIDA.match(v):
                avvisi.append(
                    f"[{eti}] data '{v}' non è nel formato 'mmm AAAA' italiano "
                    f"(mesi validi: {', '.join(sorted(MESI_IT))})")
        aperture = []
        for punto in esp.get("punti") or []:
            if APERTURE_VIETATE.match(punto):
                errori.append(f"[{eti}] bullet in prima persona o passivo: «{punto[:60]}...»")
            if PARTICIPI_ACCORDATI.match(punto):
                errori.append(
                    f"[{eti}] participio accordato in apertura, va al maschile "
                    f"singolare: «{punto[:60]}...»")
            prima = punto.strip().split(" ")[0].lower().strip(",.;:")
            if prima in aperture:
                avvisi.append(f"[{eti}] il verbo «{prima}» apre più di un bullet nello stesso ruolo")
            aperture.append(prima)

    for lingua in cv.get("lingue") or []:
        liv = (lingua.get("livello") or "").lower()
        if liv and not re.search(r"\b([abc][12]|madrelingua|bilingue)\b", liv):
            avvisi.append(
                f"lingua '{lingua.get('lingua')}': livello '{lingua.get('livello')}' "
                f"non è un livello QCER (A1-C2) né «madrelingua»")

    return errori, avvisi


# --- Primitive di formattazione ------------------------------------------------

def imposta_base(doc):
    stile = doc.styles["Normal"]
    stile.font.name = FONT
    stile.font.size = Pt(CORPO_PT)
    stile.font.color.rgb = NERO
    rpr = stile.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    pf = stile.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.06

    for sezione in doc.sections:
        sezione.top_margin = Cm(MARGINE_CM)
        sezione.bottom_margin = Cm(MARGINE_CM)
        sezione.left_margin = Cm(MARGINE_CM)
        sezione.right_margin = Cm(MARGINE_CM)


def para(doc, testo="", *, size=CORPO_PT, bold=False, italic=False,
         colore=NERO, prima=0, dopo=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(prima)
    p.paragraph_format.space_after = Pt(dopo)
    if align is not None:
        p.alignment = align
    if testo:
        run = p.add_run(testo)
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = colore
    return p


def riga_mista(doc, pezzi, *, prima=0, dopo=0, size=CORPO_PT):
    """pezzi: lista di (testo, bold, italic, colore)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(prima)
    p.paragraph_format.space_after = Pt(dopo)
    for testo, bold, italic, colore in pezzi:
        run = p.add_run(testo)
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = colore
    return p


def filetto(p):
    """Bordo inferiore del paragrafo: è una proprietà di paragrafo, non una tabella,
    quindi non interferisce con il parsing."""
    ppr = p._p.get_or_add_pPr()
    bordi = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "999999")
    bordi.append(bottom)
    ppr.append(bordi)


def titolo_sezione(doc, testo):
    p = para(doc, testo.upper(), size=TITOLO_SEZIONE_PT, bold=True,
             prima=11, dopo=4)
    p.runs[0].font.color.rgb = NERO
    filetto(p)
    return p


def bullet(doc, testo):
    p = doc.add_paragraph(testo, style="List Bullet")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.55)
    p.paragraph_format.first_line_indent = Cm(-0.3)
    for run in p.runs:
        run.font.size = Pt(CORPO_PT)
        run.font.name = FONT
    return p


def periodo(dal, al):
    dal, al = (dal or "").strip(), (al or "").strip()
    if dal and al:
        return f"{dal} – {al}"
    return dal or al


# --- Sezioni -------------------------------------------------------------------

def sezione_intestazione(doc, intest):
    nome = f"{intest.get('nome', '').strip()} {intest.get('cognome', '').strip()}".strip()
    para(doc, nome, size=NOME_PT, bold=True, dopo=1)

    if intest.get("titolo"):
        para(doc, intest["titolo"], size=CONTATTI_PT + 0.5, colore=GRIGIO, dopo=3)

    contatti = [intest.get(k) for k in ("citta", "telefono", "email")]
    contatti = [c.strip() for c in contatti if c and c.strip()]
    if contatti:
        para(doc, "  ·  ".join(contatti), size=CONTATTI_PT, dopo=1)

    link = [intest.get(k) for k in ("linkedin", "github", "sito")]
    link = [l.strip() for l in link if l and l.strip()]
    if link:
        para(doc, "  ·  ".join(link), size=CONTATTI_PT, dopo=1)

    if intest.get("extra"):
        para(doc, intest["extra"].strip(), size=CONTATTI_PT, colore=GRIGIO, dopo=1)


def sezione_profilo(doc, cv):
    if not cv.get("profilo"):
        return
    titolo_sezione(doc, TITOLI_SEZIONE["profilo"])
    para(doc, cv["profilo"].strip(), dopo=2,
         align=WD_ALIGN_PARAGRAPH.LEFT)


def sezione_competenze(doc, cv):
    gruppi = [g for g in (cv.get("competenze") or []) if g.get("voci")]
    if not gruppi:
        return
    titolo_sezione(doc, TITOLI_SEZIONE["competenze"])
    for g in gruppi:
        cat = (g.get("categoria") or "").strip()
        voci = ", ".join(v.strip() for v in g["voci"] if v and v.strip())
        pezzi = []
        if cat:
            pezzi.append((f"{cat}: ", True, False, NERO))
        pezzi.append((voci, False, False, NERO))
        riga_mista(doc, pezzi, dopo=2)


def sezione_esperienze(doc, cv):
    voci = cv.get("esperienze") or []
    if not voci:
        return
    titolo_sezione(doc, TITOLI_SEZIONE["esperienze"])
    for i, esp in enumerate(voci):
        pezzi = [(esp.get("ruolo", "").strip(), True, False, NERO)]
        coda = [esp.get("azienda", "").strip(), esp.get("citta", "").strip()]
        coda = [c for c in coda if c]
        if coda:
            pezzi.append(("  ·  " + "  ·  ".join(coda), False, False, NERO))
        p = periodo(esp.get("dal"), esp.get("al"))
        if p:
            pezzi.append((f"  ·  {p}", False, True, GRIGIO))
        riga_mista(doc, pezzi, prima=6 if i else 3, dopo=1)

        if esp.get("contesto"):
            para(doc, esp["contesto"].strip(), size=CORPO_PT - 0.5,
                 italic=True, colore=GRIGIO, dopo=2)

        for punto in esp.get("punti") or []:
            bullet(doc, punto.strip())


def sezione_formazione(doc, cv):
    voci = cv.get("formazione") or []
    if not voci:
        return
    titolo_sezione(doc, TITOLI_SEZIONE["formazione"])
    for i, f in enumerate(voci):
        pezzi = [(f.get("titolo", "").strip(), True, False, NERO)]
        if f.get("votazione"):
            pezzi.append((f"  —  {f['votazione'].strip()}", False, False, NERO))
        riga_mista(doc, pezzi, prima=5 if i else 3, dopo=0)

        coda = [f.get("istituto", "").strip(), f.get("citta", "").strip()]
        coda = [c for c in coda if c]
        p = periodo(f.get("dal"), f.get("al"))
        if p:
            coda.append(p)
        if coda:
            para(doc, "  ·  ".join(coda), size=CORPO_PT - 0.5,
                 colore=GRIGIO, dopo=1)
        for nota in f.get("note") or []:
            para(doc, nota.strip(), size=CORPO_PT - 0.5, dopo=1)


def sezione_certificazioni(doc, cv):
    voci = cv.get("certificazioni") or []
    if not voci:
        return
    titolo_sezione(doc, TITOLI_SEZIONE["certificazioni"])
    for c in voci:
        pezzi = [(c.get("nome", "").strip(), True, False, NERO)]
        coda = [c.get("ente", "").strip(), c.get("data", "").strip()]
        coda = [x for x in coda if x]
        if c.get("scadenza"):
            coda.append(f"scad. {c['scadenza'].strip()}")
        if c.get("codice"):
            coda.append(f"ID {c['codice'].strip()}")
        if coda:
            pezzi.append(("  ·  " + "  ·  ".join(coda), False, False, GRIGIO))
        riga_mista(doc, pezzi, dopo=2)


def sezione_lingue(doc, cv):
    voci = cv.get("lingue") or []
    if not voci:
        return
    titolo_sezione(doc, TITOLI_SEZIONE["lingue"])
    for l in voci:
        pezzi = [(f"{l.get('lingua', '').strip()}", True, False, NERO),
                 (f" — {l.get('livello', '').strip()}", False, False, NERO)]
        if l.get("certificazione"):
            pezzi.append((f"  ·  {l['certificazione'].strip()}",
                          False, False, GRIGIO))
        riga_mista(doc, pezzi, dopo=1)


def sezione_progetti(doc, cv):
    voci = cv.get("progetti") or []
    if not voci:
        return
    titolo_sezione(doc, TITOLI_SEZIONE["progetti"])
    for i, pr in enumerate(voci):
        pezzi = [(pr.get("nome", "").strip(), True, False, NERO)]
        coda = [pr.get("contesto", "").strip()]
        p = periodo(pr.get("dal"), pr.get("al"))
        if p:
            coda.append(p)
        if pr.get("url"):
            coda.append(pr["url"].strip())
        coda = [c for c in coda if c]
        if coda:
            pezzi.append(("  ·  " + "  ·  ".join(coda), False, False, GRIGIO))
        riga_mista(doc, pezzi, prima=5 if i else 3, dopo=1)
        for punto in pr.get("punti") or []:
            bullet(doc, punto.strip())


def sezione_altro(doc, cv):
    blocchi = cv.get("altro") or []
    if not blocchi:
        return
    for b in blocchi:
        titolo_sezione(doc, (b.get("titolo") or TITOLI_SEZIONE["altro"]))
        for v in b.get("voci") or []:
            bullet(doc, v.strip())


def sezione_privacy(doc, cv):
    testo = cv.get("privacy", PRIVACY_PREDEFINITA)
    if testo is None:
        testo = PRIVACY_PREDEFINITA
    testo = testo.strip()
    if not testo:
        return
    para(doc, testo, size=PRIVACY_PT, colore=GRIGIO, prima=12)


RENDERER = {
    "profilo": sezione_profilo,
    "competenze": sezione_competenze,
    "esperienze": sezione_esperienze,
    "formazione": sezione_formazione,
    "certificazioni": sezione_certificazioni,
    "lingue": sezione_lingue,
    "progetti": sezione_progetti,
    "altro": sezione_altro,
}


# --- Orchestrazione ------------------------------------------------------------

def costruisci(cv):
    doc = Document()
    imposta_base(doc)
    sezione_intestazione(doc, cv.get("intestazione") or {})

    profilo = ((cv.get("meta") or {}).get("profilo") or "mid").lower()
    ordine = cv.get("ordine_sezioni") or ORDINE_PREDEFINITO.get(
        profilo, ORDINE_PREDEFINITO["mid"])

    for nome in ordine:
        renderer = RENDERER.get(nome)
        if renderer is None:
            print(f"  avviso: sezione sconosciuta '{nome}', ignorata", file=sys.stderr)
            continue
        renderer(doc, cv)

    sezione_privacy(doc, cv)
    return doc


def nome_file_predefinito(cv):
    i = cv.get("intestazione") or {}
    parti = [i.get("nome", ""), i.get("cognome", ""), "CV"]
    slug = "_".join(p.strip() for p in parti if p and p.strip())
    slug = re.sub(r"[^\w]+", "_", slug, flags=re.UNICODE).strip("_")
    return f"{slug or 'CV'}.docx"


def main():
    ap = argparse.ArgumentParser(
        description="Genera un CV italiano ATS-safe in DOCX da un file JSON.")
    ap.add_argument("--input", "-i", required=True, help="file JSON del CV")
    ap.add_argument("--output", "-o", help="file DOCX di destinazione")
    ap.add_argument("--check", action="store_true",
                    help="esegue solo la validazione, senza generare il DOCX")
    ap.add_argument("--force", action="store_true",
                    help="genera il DOCX anche in presenza di errori bloccanti")
    args = ap.parse_args()

    try:
        with open(args.input, encoding="utf-8") as fh:
            cv = json.load(fh)
    except json.JSONDecodeError as exc:
        sys.exit(f"JSON non valido in {args.input}: {exc}")
    except OSError as exc:
        sys.exit(f"Impossibile leggere {args.input}: {exc}")

    errori, avvisi = valida(cv)
    for a in avvisi:
        print(f"  AVVISO  {a}")
    for e in errori:
        print(f"  ERRORE  {e}", file=sys.stderr)

    if args.check:
        if errori:
            sys.exit(f"\nValidazione fallita: {len(errori)} errore/i, {len(avvisi)} avviso/i.")
        print(f"\nValidazione superata ({len(avvisi)} avviso/i).")
        return

    if errori and not args.force:
        sys.exit(
            f"\n{len(errori)} errore/i bloccante/i. "
            "Correggi il JSON, oppure rilancia con --force se sono falsi positivi.")

    uscita = args.output or nome_file_predefinito(cv)
    try:
        costruisci(cv).save(uscita)
    except OSError as exc:
        sys.exit(f"Impossibile scrivere {uscita}: {exc}")
    print(f"\nCV generato: {uscita}")


if __name__ == "__main__":
    main()
